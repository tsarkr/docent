"""Excel helpers for the human annotation sheets (expert validation, answer rating)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

EXCEL_CELL_LIMIT = 32000


def add_sheet(
    wb: Workbook,
    title: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    widths: dict[str, int] | None = None,
    choices: dict[str, list[str]] | None = None,
) -> None:
    """Append a sheet; ``choices`` maps a column to its allowed dropdown values."""
    ws = wb.create_sheet(title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([str(row.get(c, ""))[:EXCEL_CELL_LIMIT] if row.get(c) is not None else "" for c in columns])
    ws.freeze_panes = "A2"
    for idx, column in enumerate(columns, 1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = (widths or {}).get(column, 16)
        if (widths or {}).get(column, 0) >= 40:
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if choices and column in choices:
            validation = DataValidation(type="list", formula1='"' + ",".join(choices[column]) + '"', allow_blank=True)
            ws.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{max(len(rows) + 1, 2000)}")


def new_workbook(instructions: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    for line in instructions:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    return wb


def read_sheet(path: Path, title: str) -> list[dict[str, str]]:
    ws = load_workbook(path, read_only=True, data_only=True)[title]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    out = []
    for values in rows:
        record = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, values)}
        if any(record.values()):
            out.append(record)
    return out
